"""Generate the Word fixture for the import acceptance run.

A manuscript written the way a person writes in Word: a title, chapter
headings, body text with italics and bold, a scene break, a block quote, a
bulleted list, comments with ranges, and tracked insertions and deletions
(python-docx has no API for tracked changes, so the w:ins / w:del elements are
written into the XML directly, the way Word writes them).

    python make_docx.py <out.docx>
"""
import copy
import sys
from datetime import datetime, timezone

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def tracked(paragraph, kind, text, author="Reviewer", revision=[100]):
    """Append a tracked insertion or deletion run to a paragraph."""
    revision[0] += 1
    wrapper = OxmlElement("w:ins" if kind == "ins" else "w:del")
    wrapper.set(qn("w:id"), str(revision[0]))
    wrapper.set(qn("w:author"), author)
    wrapper.set(qn("w:date"), datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))
    run = OxmlElement("w:r")
    node = OxmlElement("w:t" if kind == "ins" else "w:delText")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    wrapper.append(run)
    paragraph._p.append(wrapper)


def build(path):
    doc = Document()
    doc.core_properties.title = "The Salt Road"
    doc.core_properties.author = "Acceptance"
    doc.add_heading("The Salt Road", level=0)
    doc.add_heading("Chapter One: The Harbour", level=1)
    p = doc.add_paragraph("Mira tied back her ")
    p.add_run("auburn").italic = True
    p.add_run(" hair and waited for Oren, who was ")
    p.add_run("never").bold = True
    p.add_run(" on time.")
    doc.add_comment(
        runs=[p.runs[1]],
        text="Check the hair colour against chapter three.",
        author="Ida",
        initials="I",
    )
    p2 = doc.add_paragraph("The tide was out. ")
    tracked(p2, "ins", "A gull cried twice.")
    tracked(p2, "del", " The sea was loud.")
    doc.add_paragraph("* * *")
    q = doc.add_paragraph("Salt keeps what the sea forgets.")
    q.style = doc.styles["Quote"]
    doc.add_paragraph("Rope", style="List Bullet")
    doc.add_paragraph("Lantern", style="List Bullet")
    doc.add_heading("Chapter Two: The Ledger", level=1)
    p3 = doc.add_paragraph("Mira shook out her ")
    p3.add_run("black").italic = True
    p3.add_run(" hair and opened the ledger.")
    doc.add_comment(
        runs=[p3.runs[1]],
        text="Black? Earlier she was auburn.",
        author="Ida",
        initials="I",
    )
    doc.add_paragraph("The ledger was empty but for one name.")
    doc.add_heading("Chapter Three: The Lantern", level=1)
    doc.add_paragraph("Nobody climbed the lantern stairs after dark.")
    doc.save(path)


if __name__ == "__main__":
    build(sys.argv[1])
    print(sys.argv[1])
